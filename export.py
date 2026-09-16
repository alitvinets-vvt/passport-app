"""
Експорт результату розрахунку Планового Паспорта в xlsx.

Чиста презентаційна функція — не рахує нічого сама, лише переносить
1:1 значення, які вже повернула calculator.calculate().
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BORDO = "7A2331"
CREAM = "F1E9DC"

SECTION_FONT = Font(bold=True, color=BORDO, size=12)
SECTION_FILL = PatternFill(start_color=CREAM, end_color=CREAM, fill_type="solid")
HEADER_FONT = Font(bold=True, color=BORDO)
TOTAL_FONT = Font(bold=True)
MONEY_FORMAT = "#,##0.00"
INT_FORMAT = "#,##0"

ROW_LABELS = {
    "аванс": "Аванс за текст",
    "переклад": "Переклад",
    "редагування": "Редагування",
    "коректура": "Коректура",
    "обкладинка": "Обкладинка (дизайн)",
    "ефекти": "Ефекти обкладинки",
}


def _write_section_title(ws, row, title, span=4):
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = SECTION_FONT
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = SECTION_FILL
    return row + 1


def _display(value, suffix=""):
    if value is None:
        return "—"
    return f"{value}{suffix}"


def export_to_xlsx(inputs: dict, result: dict) -> io.BytesIO:
    """Формує xlsx: шапка (індекс/назва проєкту) і три секції —
    параметри книги, редакційні витрати (Блок 1) і друк + підсумок
    (Блок 2 + РРЦ + точка беззбитковості) — на одному аркуші, значення
    1:1 з того, що показано в UI.

    inputs["project_index"]/["project_name"] — ручна ідентифікація
    проєкту для файлу експорту, не бере участі в calculate().
    result["breakeven"] — з calculate(), None якщо не порахована
    (тоді відповідні клітинки лишаються порожніми)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Плановий Паспорт"

    row = 1

    # ---- Шапка: ідентифікація проєкту ----
    ws.cell(row=row, column=1, value="Індекс проєкту:")
    ws.cell(row=row, column=2, value=inputs.get("project_index", ""))
    row += 1
    ws.cell(row=row, column=1, value="Назва проєкту:")
    ws.cell(row=row, column=2, value=inputs.get("project_name", ""))
    row += 2

    # ---- Секція 1: Параметри книги ----
    row = _write_section_title(ws, row, "Параметри книги")
    book_params = [
        ("Формат", inputs.get("format")),
        ("Зірковість", inputs.get("zirka")),
        ("Складність", inputs.get("complexity")),
        ("Кількість знаків (оригінал)", inputs.get("znaky")),
        ("Перекладна книга", "Так" if inputs.get("perekladna") else "Ні"),
        ("Наклад", inputs.get("naklad")),
        ("Множник сторінковості", inputs.get("storinkovist_multiplier", 1.0)),
        ("Колірність блоку", inputs.get("color_mode")),
        ("Ефекти обкладинки", inputs.get("effect")),
        ("Кольоровий зріз", "Так" if inputs.get("has_zriz") else "Ні"),
    ]
    for label, value in book_params:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.number_format = "0.00" if label == "Множник сторінковості" else INT_FORMAT
        row += 1

    row += 1

    # ---- Секція 2: Редакційні витрати (Блок 1) ----
    row = _write_section_title(ws, row, "Редакційні витрати")
    headers = ["Стаття", "Чистими, грн", "Тіло, грн", "ЄСВ, грн", "Разом, грн"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
    row += 1

    for key, label in ROW_LABELS.items():
        r = result["rows"][key]
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=r["chysto"]).number_format = MONEY_FORMAT
        ws.cell(row=row, column=3, value=r["tilo"]).number_format = MONEY_FORMAT
        ws.cell(row=row, column=4, value=r["esv"]).number_format = MONEY_FORMAT
        ws.cell(row=row, column=5, value=r["razom"]).number_format = MONEY_FORMAT
        row += 1

    inshi_row = result["inshi_row"]
    ws.cell(row=row, column=1, value="Інші витрати (12%)")
    ws.cell(row=row, column=2, value=inshi_row["chysto"]).number_format = MONEY_FORMAT
    ws.cell(row=row, column=3, value=inshi_row["tilo"]).number_format = MONEY_FORMAT
    ws.cell(row=row, column=4, value=inshi_row["esv"]).number_format = MONEY_FORMAT
    ws.cell(row=row, column=5, value=inshi_row["razom"]).number_format = MONEY_FORMAT
    row += 1

    ws.cell(row=row, column=1, value="ОРИГІНАЛ-МАКЕТ").font = TOTAL_FONT
    total_cell = ws.cell(row=row, column=5, value=result["oryhinal_maket"])
    total_cell.number_format = MONEY_FORMAT
    total_cell.font = TOTAL_FONT
    row += 2

    # ---- Секція 3: Друк і підсумок (Блок 2 + фінал) ----
    row = _write_section_title(ws, row, "Друк і підсумок")
    block2 = result.get("block2")

    if block2 is None:
        ws.cell(row=row, column=1, value="Друк не порахований — бракує даних у майстер-таблиці.")
        row += 1
    else:
        print_rows = [
            ("Сторінки", block2["storinky"], "0.0"),
            ("Зошити", block2["zoshytiv"], INT_FORMAT),
            ("Тир", f"{block2['tier']['tier']} (k={block2['tier']['k']})", None),
            ("Блок (друк)", block2["blok"], MONEY_FORMAT),
            ("Обкладинка (друк)", block2["obkladynka_dr"], MONEY_FORMAT),
            ("Кольоровий зріз (друк)", block2["zriz"], MONEY_FORMAT),
        ]
        for label, value, fmt in print_rows:
            ws.cell(row=row, column=1, value=label)
            cell = ws.cell(row=row, column=2, value=value)
            if fmt:
                cell.number_format = fmt
            row += 1

        ws.cell(row=row, column=1, value="Друк за 1 прим., разом").font = TOTAL_FONT
        druk_cell = ws.cell(row=row, column=2, value=block2["razom"])
        druk_cell.number_format = MONEY_FORMAT
        druk_cell.font = TOTAL_FONT
        row += 2

    ws.cell(row=row, column=1, value="РРЦ").font = SECTION_FONT
    rrc_cell = ws.cell(row=row, column=2, value=result.get("rrc"))
    if result.get("rrc") is not None:
        rrc_cell.number_format = INT_FORMAT
    rrc_cell.font = SECTION_FONT
    row += 2

    breakeven = result.get("breakeven")
    ws.cell(row=row, column=1, value="Знижка рітейлу, %")
    ws.cell(
        row=row, column=2,
        value=breakeven["retail_discount"] * 100 if breakeven else None,
    ).number_format = "0"
    row += 1

    ws.cell(row=row, column=1, value="Точка беззбитковості, прим.").font = TOTAL_FONT
    units_cell = ws.cell(row=row, column=2, value=breakeven["units"] if breakeven else None)
    units_cell.number_format = INT_FORMAT
    units_cell.font = TOTAL_FONT
    row += 1

    ws.cell(row=row, column=1, value="Точка беззбитковості, % тиражу")
    ws.cell(
        row=row, column=2,
        value=breakeven["percent_of_run"] if breakeven else None,
    ).number_format = "0.0"

    # ---- Ширина колонок ----
    widths = [28, 16, 16, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
